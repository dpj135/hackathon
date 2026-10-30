#!/usr/bin/env python3
"""Render submission text to DOCX and offline HTML; requires python-docx."""
import html
import re
from datetime import datetime, timezone
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parents[1]
CSS='''body{font-family:system-ui,"Microsoft YaHei",sans-serif;color:#183047;background:#f4f7fa;line-height:1.85;margin:0}main{max-width:1000px;margin:28px auto;padding:34px;background:white}h1{font-size:30px}h2{margin-top:32px;color:#1647d8;font-size:22px}h3{font-size:18px}p,li{font-size:16px}table{border-collapse:collapse;width:100%;font-size:14px;margin:18px 0}th,td{padding:10px;border:1px solid #dce8ff;text-align:left;vertical-align:top}th{background:#eaf0ff}pre{background:#eef3f8;padding:16px;overflow:auto;line-height:1.6}code{font-family:monospace;overflow-wrap:anywhere}a{color:#1647d8}.table{overflow:auto}@media(max-width:700px){main{margin:0;padding:18px}h1{font-size:25px}}@media print{body{background:white}main{max-width:none;margin:0;padding:0}h2{break-after:avoid}}'''

def inline(s):
    s=html.escape(s)
    s=re.sub(r'`([^`]+)`',r'<code>\1</code>',s)
    s=re.sub(r'\*\*([^*]+)\*\*',r'<strong>\1</strong>',s)
    return re.sub(r'\[([^]]+)\]\(([^)]+)\)',r'<a href="\2">\1</a>',s)

def plain(s):
    return re.sub(r'\[([^]]+)\]\(([^)]+)\)',r'\1（\2）',s).replace('**','').replace('`','')

def blocks(text):
    lines=text.splitlines();i=0
    while i<len(lines):
        line=lines[i];i+=1
        if not line.strip():continue
        if line.startswith('```'):
            code=[]
            while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
            i+=1;yield 'code','\n'.join(code);continue
        if line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].startswith('|'):rows.append(lines[i]);i+=1
            yield 'table',[[s.strip() for s in row.strip('|').split('|')] for row in rows if not re.fullmatch(r'[| :\-]+',row)];continue
        if line.startswith('#'):
            depth=len(line)-len(line.lstrip('#'));yield 'heading',(min(depth,3),line.lstrip('#').strip());continue
        if line.startswith('- '):yield 'bullet',line[2:];continue
        if re.match(r'^\d+\. ',line):yield 'number',re.sub(r'^\d+\. ','',line);continue
        yield 'paragraph',line

def render(title,files,dest):
    doc=Document();doc.core_properties.title=title;doc.core_properties.author='';doc.core_properties.last_modified_by='';doc.core_properties.created=doc.core_properties.modified=datetime.now(timezone.utc);doc.core_properties.comments=''
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        style=doc.styles[name];style.font.name='Microsoft YaHei';style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Microsoft YaHei')
    doc.styles['Normal'].font.size=Pt(10.5);doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    for section in doc.sections:
        section.top_margin=section.bottom_margin=Inches(.7);section.left_margin=section.right_margin=Inches(.75)
        section.header.paragraphs[0].text='NDS 多路径IO分配聚合与缓存提升调度'
        footer=section.footer.paragraphs[0];footer.alignment=2
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
    content=[]
    if len(files)>1:
        doc.add_heading(title,0)
        doc.add_paragraph('观察窗口调度与复用感知缓存',style='Subtitle')
        doc.add_paragraph('同一份C++17策略，对应源码、可执行程序、验证证据和答辩材料。')
        doc.add_paragraph('团队 / 作者：________________')
        content.append('<h1>'+html.escape(title)+'</h1><p>观察窗口调度与复用感知缓存</p>')
    for index,file in enumerate(files):
        if index:doc.add_page_break()
        number=0
        for kind,value in blocks((root/file).read_text()):
            number=number+1 if kind=='number' else 0
            if kind=='heading':
                level,text=value;doc.add_heading(plain(text),level);content.append(f'<h{level}>{inline(text)}</h{level}>')
            elif kind=='table':
                table=doc.add_table(rows=0,cols=len(value[0]));table.style='Light Shading Accent 1';content.append('<div class="table"><table>')
                for ri,row in enumerate(value):
                    for cell,text in zip(table.add_row().cells,row):
                        cell.text=plain(text)
                        for p in cell.paragraphs:
                            for run in p.runs:run.font.size=Pt(9)
                    tag='th' if ri==0 else 'td';content.append('<tr>'+''.join(f'<{tag}>{inline(x)}</{tag}>' for x in row)+'</tr>')
                content.append('</table></div>')
            elif kind=='code':
                for line in value.splitlines():
                    run=doc.add_paragraph().add_run(line);run.font.name='Courier New';run.font.size=Pt(8)
                content.append('<pre><code>'+html.escape(value)+'</code></pre>')
            else:
                style='List Bullet' if kind=='bullet' else 'List Number' if kind=='number' else None
                doc.add_paragraph(plain(value),style=style)
                if kind in ['bullet','number']:
                    tag='ul' if kind=='bullet' else 'ol';content.append((f'<ol start="{number}">' if kind=='number' else '<ul>')+f'<li>{inline(value)}</li></{tag}>')
                else:content.append('<p>'+inline(value)+'</p>')
    path=root/dest;doc.save(path.with_suffix('.docx'))
    page='<!doctype html><html lang="zh-CN"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>'+html.escape(title)+'</title><style>'+CSS+'</style></head><body><main>'+''.join(content)+'</main></body></html>'
    path.with_suffix('.html').write_text(page,encoding='utf-8');print(path)

render('从零认识我们的 NDS 调度方案',['00_快速入门/方案介绍.md'],'00_快速入门/方案介绍')
render('NDS 解题方案与验证报告',['01_题目与依据/材料来源.md','02_算法方案/算法设计.md','02_算法方案/通用性与边界条件.md','05_验证结果/验证报告.md','04_测试复现/测试操作手册.md'],'02_算法方案/解题方案与验证报告')

render('验证报告',['05_验证结果/验证报告.md'],'05_验证结果/验证报告')
